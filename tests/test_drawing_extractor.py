"""DrawingML抽出モジュールのユニットテスト。"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import openpyxl
import pytest

from excel_to_markdown.drawing.extractor import extract_diagrams, extract_sheet_drawing_map
from excel_to_markdown.models import DiagramConnector, DiagramShape

SAMPLE_FLOWCHART = Path(__file__).parent / "e2e/fixtures/sample_flowchart.xlsx"


# ---------------------------------------------------------------------------
# フィクスチャ: drawingなしのxlsx
# ---------------------------------------------------------------------------

@pytest.fixture()
def xlsx_no_drawing(tmp_path: Path) -> Path:
    """drawingを持たないxlsxを返す。"""
    path = tmp_path / "no_drawing.xlsx"
    wb = openpyxl.Workbook()
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# drawingなしのケース
# ---------------------------------------------------------------------------

class TestNoDiagram:
    def test_returns_empty_list(self, xlsx_no_drawing: Path) -> None:
        results = extract_diagrams(xlsx_no_drawing)
        assert results == []


# ---------------------------------------------------------------------------
# サンプルフローチャートの抽出テスト
# ---------------------------------------------------------------------------

class TestSampleFlowchart:
    @pytest.fixture(autouse=True)
    def setup(self) -> None:
        if not SAMPLE_FLOWCHART.exists():
            pytest.skip(f"サンプルファイルが見つかりません: {SAMPLE_FLOWCHART}")
        results = extract_diagrams(SAMPLE_FLOWCHART)
        assert len(results) == 1
        self.shapes, self.connectors = results[0]

    # ---- 図形の数と内容 ----

    def test_shape_count(self) -> None:
        assert len(self.shapes) == 7

    def test_shape_types(self) -> None:
        type_counts: dict[str, int] = {}
        for s in self.shapes:
            type_counts[s.shape_type] = type_counts.get(s.shape_type, 0) + 1
        assert type_counts.get("flowChartTerminator", 0) == 2
        assert type_counts.get("flowChartProcess", 0) == 4
        assert type_counts.get("flowChartDecision", 0) == 1

    def test_shape_texts(self) -> None:
        texts = {s.text for s in self.shapes}
        assert "開始" in texts
        assert "完了" in texts
        assert "受注受付" in texts
        assert "在庫確認?" in texts
        assert "出荷処理" in texts
        assert "発注処理" in texts
        assert "入荷待ち" in texts

    def test_shape_ids_unique(self) -> None:
        ids = [s.shape_id for s in self.shapes]
        assert len(ids) == len(set(ids))

    # ---- コネクタの数と内容 ----

    def test_connector_count(self) -> None:
        assert len(self.connectors) == 7

    def test_connector_start_end_ids(self) -> None:
        edges = {(c.start_shape_id, c.end_shape_id) for c in self.connectors}
        shape_by_text = {s.text: s.shape_id for s in self.shapes}

        # 主要な接続を確認
        assert (shape_by_text["開始"], shape_by_text["受注受付"]) in edges
        assert (shape_by_text["受注受付"], shape_by_text["在庫確認?"]) in edges
        assert (shape_by_text["在庫確認?"], shape_by_text["出荷処理"]) in edges
        assert (shape_by_text["在庫確認?"], shape_by_text["発注処理"]) in edges
        assert (shape_by_text["出荷処理"], shape_by_text["完了"]) in edges

    def test_connector_labels(self) -> None:
        labels = {c.label for c in self.connectors if c.label}
        assert "はい" in labels
        assert "いいえ" in labels

    def test_connectors_have_both_endpoints(self) -> None:
        for c in self.connectors:
            assert c.start_shape_id is not None, f"コネクタ {c.connector_id} の始点がNone"
            assert c.end_shape_id is not None, f"コネクタ {c.connector_id} の終点がNone"

    # ---- 位置情報 ----

    def test_shape_position_non_negative(self) -> None:
        for s in self.shapes:
            assert s.left_col >= 0
            assert s.top_row >= 0
            assert s.right_col >= s.left_col
            assert s.bottom_row >= s.top_row


# ---------------------------------------------------------------------------
# DiagramShape / DiagramConnector の dataclass 仕様テスト
# ---------------------------------------------------------------------------

class TestDataclasses:
    def test_diagram_shape_immutable(self) -> None:
        shape = DiagramShape(
            shape_id=1,
            name="テスト",
            text="テキスト",
            shape_type="rect",
            left_col=0,
            top_row=0,
            right_col=2,
            bottom_row=2,
        )
        with pytest.raises(Exception):
            shape.shape_id = 99  # type: ignore[misc]

    def test_diagram_connector_immutable(self) -> None:
        connector = DiagramConnector(
            connector_id=1,
            name="矢印",
            start_shape_id=2,
            end_shape_id=3,
            label="",
        )
        with pytest.raises(Exception):
            connector.connector_id = 99  # type: ignore[misc]

    def test_connector_none_endpoints(self) -> None:
        """未接続コネクタ（start/end = None）が正常に作成できる。"""
        connector = DiagramConnector(
            connector_id=1,
            name="未接続",
            start_shape_id=None,
            end_shape_id=None,
            label="",
        )
        assert connector.start_shape_id is None
        assert connector.end_shape_id is None


# ---------------------------------------------------------------------------
# extract_sheet_drawing_map のテスト
# ---------------------------------------------------------------------------

SAMPLE_1 = Path(__file__).parent / "e2e/fixtures/1.xlsx"
SAMPLE_GYOUMU = Path(__file__).parent / "e2e/fixtures/gyoumuflow_answer.xlsx"


class TestExtractSheetDrawingMap:
    def test_no_drawing_returns_empty(self, xlsx_no_drawing: Path) -> None:
        result = extract_sheet_drawing_map(xlsx_no_drawing)
        assert result == {}

    def test_sample_flowchart_maps_sheet(self) -> None:
        if not SAMPLE_FLOWCHART.exists():
            pytest.skip("サンプルファイルが見つかりません")
        result = extract_sheet_drawing_map(SAMPLE_FLOWCHART)
        assert len(result) == 1
        sheet_name = next(iter(result))
        shapes, connectors = result[sheet_name]
        assert len(shapes) == 7
        assert len(connectors) == 7

    def test_1xlsx_sheet_names(self) -> None:
        if not SAMPLE_1.exists():
            pytest.skip("1.xlsx が見つかりません")
        result = extract_sheet_drawing_map(SAMPLE_1)
        # drawing を持つシートが正しく検出される
        assert "画面遷移" in result
        assert "トップ画面" in result
        # drawing のないシートはマップに含まれない
        assert "共通要件" not in result

    def test_1xlsx_画面遷移_shapes(self) -> None:
        if not SAMPLE_1.exists():
            pytest.skip("1.xlsx が見つかりません")
        result = extract_sheet_drawing_map(SAMPLE_1)
        shapes, connectors = result["画面遷移"]
        assert len(shapes) > 0
        assert len(connectors) > 0
        # ログインフローの主要図形が含まれる
        texts = {s.text for s in shapes}
        assert "ログイン画面" in texts
        assert "トップ画面" in texts

    def test_gyoumu_sheet_names(self) -> None:
        if not SAMPLE_GYOUMU.exists():
            pytest.skip("gyoumuflow_answer.xlsx が見つかりません")
        result = extract_sheet_drawing_map(SAMPLE_GYOUMU)
        assert "単一業務" in result
        assert "複数部署" in result
        assert "はじめに" not in result

    def test_gyoumu_複数部署_flow(self) -> None:
        if not SAMPLE_GYOUMU.exists():
            pytest.skip("gyoumuflow_answer.xlsx が見つかりません")
        result = extract_sheet_drawing_map(SAMPLE_GYOUMU)
        shapes, connectors = result["複数部署"]
        texts = {s.text for s in shapes}
        assert "受付用紙に記入する" in texts
        assert "工事を行う" in texts
        # コネクタが存在する（接続あり）
        assert len(connectors) > 0
